# -*- coding: utf-8 -*-
"""Lista las siglas usadas en el cuerpo de una presentación y las que faltan en su glosario."""
import re, sys
from pptx import Presentation
IGNORAR = {"SOLO", "DATOS", "LLAVE", "LLAVES", "MAESTRA", "BEGIN", "CERTIFICATE", "F1", "F2", "FALSE", "FINAL", "INTERMEDIA", "KEY", "PRIVATE", "REQUEST", "BIEN", "FIRMAR", "MAL", "SUMA", "VERIFICAR", "CON", "SIN", "ESE", "H12", "H34", "A7F3", "CIFRADA", "CIFRADO", "HACE", "NADA", "REALMENTE", "SALDO", "UN", "VALOR", "YA", "AESGCM", "NO", "SI", "C1", "C2", "M1", "M2", "COOPABURRA", "FRRSDEXUUD", "SIO0010", "CRC", "MIFARE", "II", "III", "IV", "VI", "XX", "XIX", "XXI", "UTF", "OK", "MAR", "EL", "LA", "DE", "EN", "QU", "ES",
           "CIFRA", "PAGAR", "TODO", "VIERNES", "RERX", "PTLN", "GDIS", "AOVE", "AOEX", "CA", "CN", "GMT"}
pr = Presentation(sys.argv[1])
usadas, glosario = {}, set()
for i, s in enumerate(pr.slides, 1):
    cajas = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    es_glosario = any("SIGLAS DE LA SESIÓN" in c for c in cajas)
    for c in cajas:
        if es_glosario:
            if re.fullmatch(r"[A-Z0-9][A-Za-z0-9\-/ …\.·]{0,44}", c.strip()):
                for parte in re.split(r"\s*[/·]\s*|\s*…\s*", c.strip()):
                    glosario.add(parte.strip())
                    for pieza in parte.split():
                        glosario.add(pieza.strip())
            continue
        if c.strip().upper() == c.strip():          # títulos y antetítulos: todo en mayúsculas
            continue
        for m in re.finditer(r"\b([A-Z][A-Z0-9]{1,}(?:-[A-Z0-9]+)?)\b", c):
            t = m.group(1)
            if t in IGNORAR or re.fullmatch(r"H\d", t) or t.isdigit():
                continue
            usadas.setdefault(t, set()).add(i)
cubierta = lambda t: t in glosario or t.split("-")[0] in glosario or any(t.startswith(g) for g in glosario if len(g) > 2)
faltan = {t: v for t, v in usadas.items() if not cubierta(t)}
print("Glosario:", ", ".join(sorted(glosario)) or "—")
print("Faltan en el glosario:")
for t in sorted(faltan):
    print(f"   {t:14s} diapositivas {sorted(faltan[t])[:8]}")
