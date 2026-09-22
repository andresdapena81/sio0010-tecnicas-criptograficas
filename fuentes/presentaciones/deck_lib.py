# -*- coding: utf-8 -*-
"""Sistema de diseño replicado de SIO0012-E1-Gobierno-de-la-seguridad.pptx"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------- tokens ----------
NEGRO   = RGBColor(0x00, 0x00, 0x00)
NARANJA = RGBColor(0xF3, 0x64, 0x06)
GRIS    = RGBColor(0x5C, 0x5C, 0x5C)
GRIS_PIE= RGBColor(0x70, 0x70, 0x70)
GRIS_LBL= RGBColor(0x3A, 0x3A, 0x3A)
TINTA   = RGBColor(0x2A, 0x2A, 0x2A)
LINEA   = RGBColor(0xC9, 0xC9, 0xC9)
BLANCO  = RGBColor(0xFF, 0xFF, 0xFF)

M, W = 0.62, 12.09          # margen y ancho útil
TITULAR, MONO, CUERPO = "Arial Black", "Consolas", "Arial"
PIE_TXT = "SIO0010 · TÉCNICAS CRIPTOGRÁFICAS · SESIÓN 1"


def nueva(logo_grande, logo_pie):
    p = Presentation()
    p.slide_width, p.slide_height = Emu(12192000), Emu(6858000)
    p._logo_g, p._logo_p = logo_grande, logo_pie
    return p


def _marcas(t, base):
    """Divide un texto con <b>…</b> en trozos (texto, negrita)."""
    if "<b>" not in t:
        return [(t, base)]
    out, resto = [], t
    while "<b>" in resto:
        antes, resto = resto.split("<b>", 1)
        if antes: out.append((antes, base))
        dentro, resto = resto.split("</b>", 1) if "</b>" in resto else (resto, "")
        if dentro: out.append((dentro, True))
    if resto: out.append((resto, base))
    return out


def _tx(s, x, y, w, h, texto, fuente, tam, color, bold=False, spc=None,
        align=PP_ALIGN.LEFT, interlineado=None, espacio_post=0):
    """Cuadro de texto. `texto` puede ser str o lista de (str, dict) para runs mixtos."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    parrafos = texto if isinstance(texto, list) else [texto]
    for i, item in enumerate(parrafos):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        if interlineado: par.line_spacing = interlineado
        par.space_after = Pt(espacio_post)
        runs = item if isinstance(item, list) else [(item, {})]
        for t, op in runs:
            for trozo, negrita in _marcas(t, op.get("bold", bold)):
                r = par.add_run(); r.text = trozo
                f = r.font
                f.name = op.get("fuente", fuente)
                f.size = Pt(op.get("tam", tam))
                f.bold = negrita
                f.color.rgb = op.get("color", color)
                e = op.get("spc", spc)
                if e: r.font._rPr.set("spc", str(int(e)))
    return tb


def _rect(s, x, y, w, h, relleno=None, borde=None, grosor=1.4):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if relleno is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = relleno
    if borde is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = borde; sh.line.width = Pt(grosor)
    return sh


def _pie(p, s, num):
    _rect(s, M, 6.80, W, 0.01, relleno=LINEA)
    s.shapes.add_picture(p._logo_p, Inches(M), Inches(6.93), Inches(0.62), Inches(0.27))
    _tx(s, 1.47, 6.94, 9.20, 0.25, PIE_TXT, MONO, 7, GRIS_PIE, bold=True, spc=150)
    _tx(s, 11.51, 6.86, 1.20, 0.42, f"{num:02d}", TITULAR, 15, NEGRO, align=PP_ALIGN.RIGHT)


def portada(p, kicker, titulo_lineas, pregunta, creditos):
    s = p.slides.add_slide(p.slide_layouts[6])
    s.shapes.add_picture(p._logo_g, Inches(M), Inches(0.55), Inches(1.55), Inches(0.67))
    _rect(s, M, 1.95, W, 0.06, relleno=NEGRO)
    _tx(s, M, 2.16, W, 0.26, kicker, MONO, 8.5, GRIS, bold=True, spc=220)
    _tx(s, M, 2.50, 10.5, 2.35, titulo_lineas, TITULAR, 40, NEGRO, interlineado=1.05)
    _rect(s, M, 5.02, 1.70, 0.06, relleno=NARANJA)
    _tx(s, M, 5.30, 11.4, 0.40, pregunta, CUERPO, 14, NEGRO, bold=True)
    _tx(s, M, 5.78, 11.4, 0.55, creditos, CUERPO, 11, GRIS, interlineado=1.25)
    _pie(p, s, 1)
    return s


def seccion(p, num_sec, titulo, bajada, num):
    s = p.slides.add_slide(p.slide_layouts[6])
    _tx(s, M, 2.30, W, 0.26, "SECCIÓN", MONO, 8.5, GRIS, bold=True, spc=240)
    _tx(s, M, 2.60, 2.0, 1.10, num_sec, TITULAR, 54, NARANJA)
    _tx(s, M, 3.85, W, 0.70, titulo, TITULAR, 30, NEGRO)
    _rect(s, M, 4.72, 1.70, 0.05, relleno=NEGRO)
    _tx(s, M, 4.95, 11.0, 0.50, bajada, CUERPO, 12, GRIS, interlineado=1.3)
    _pie(p, s, num)
    return s


def tesis(p, eyebrow, frase_lineas, apoyo, num):
    s = p.slides.add_slide(p.slide_layouts[6])
    _tx(s, M, 1.30, W, 0.26, eyebrow, MONO, 8.5, GRIS, bold=True, spc=240)
    _tx(s, M, 1.70, 11.6, 3.0, frase_lineas, TITULAR, 30, NEGRO, interlineado=1.12)
    _rect(s, M, 5.05, 1.70, 0.05, relleno=NARANJA)
    _tx(s, M, 5.32, 11.6, 1.10, apoyo, CUERPO, 12, TINTA, interlineado=1.35)
    _pie(p, s, num)
    return s


def base(p, eyebrow, titulo, num, titulo_tam=26, lineas_titulo=1):
    """Diapositiva de contenido. Devuelve (slide, y_inicial_del_contenido)."""
    s = p.slides.add_slide(p.slide_layouts[6])
    _rect(s, M, 0.40, W, 0.05, relleno=NEGRO)
    _tx(s, M, 0.56, W, 0.22, eyebrow, MONO, 8.5, GRIS, bold=True, spc=240)
    alto = 0.62 * lineas_titulo
    _tx(s, M, 0.80, W, alto, titulo, TITULAR, titulo_tam, NEGRO, interlineado=1.08)
    _pie(p, s, num)
    return s, 0.80 + alto + 0.28


def intro(s, y, texto, tam=12.5):
    """Párrafo de entrada bajo el título. Devuelve el nuevo y."""
    def _plano(x):
        if isinstance(x, str): return x
        if isinstance(x, tuple): return _plano(x[0])
        return "".join(_plano(i) for i in x)
    n = 1 + len(_plano(texto)) // 132
    _tx(s, M, y, W, 0.26 * n, texto, CUERPO, tam, TINTA, interlineado=1.35)
    return y + 0.26 * n + 0.30


def cifras(s, y, datos, alto=1.16):
    """Fila de cifras: [(valor, etiqueta_l1, etiqueta_l2, destacado_bool), ...]"""
    _rect(s, M, y, W, 0.03, relleno=NEGRO)
    y2 = y + 0.14
    ancho = W / len(datos)
    for i, (val, l1, l2, dest) in enumerate(datos):
        x = M + i * ancho
        if i:
            _rect(s, x - 0.12, y2 - 0.02, 0.01, 1.08, relleno=NEGRO)
        _tx(s, x, y2, ancho - 0.20, 0.60, val, TITULAR, 28, NARANJA if dest else NEGRO)
        _tx(s, x, y2 + 0.66, ancho - 0.20, 0.48,
            [l1, l2] if l2 else [l1], MONO, 7.3, GRIS_LBL, bold=True, spc=130, interlineado=1.25)
    _rect(s, M, y2 + 1.14, W, 0.03, relleno=NEGRO)
    return y2 + 1.14 + 0.03 + 0.34


def tarjetas(s, y, items, alto=2.06, cols=2):
    """items: [(eyebrow, [parrafos])]"""
    gap = 0.38
    ancho = (W - gap * (cols - 1)) / cols
    for i, (eb, parrafos) in enumerate(items):
        col, fila = i % cols, i // cols
        x = M + col * (ancho + gap)
        yy = y + fila * (alto + 0.30)
        _rect(s, x, yy, ancho, alto, relleno=BLANCO, borde=NEGRO, grosor=1.4)
        _tx(s, x + 0.24, yy + 0.12, ancho - 0.48, 0.22, eb, MONO, 8, NEGRO, bold=True, spc=190)
        _tx(s, x + 0.24, yy + 0.36, ancho - 0.48, alto - 0.50, parrafos,
            CUERPO, 11.5, NEGRO, interlineado=1.28, espacio_post=7)
    filas = (len(items) + cols - 1) // cols
    return y + filas * (alto + 0.30) + 0.10


def tabla(s, y, encabezados, filas, anchos, alto_fila=0.40, tam=11):
    """Tabla en el estilo del módulo: reglas negras, sin relleno."""
    total = sum(anchos)
    anchos = [a * W / total for a in anchos]
    _rect(s, M, y, W, 0.03, relleno=NEGRO)
    yy = y + 0.13
    x = M
    for i, h in enumerate(encabezados):
        _tx(s, x, yy, anchos[i] - 0.18, 0.30, h, MONO, 7.3, GRIS_LBL, bold=True, spc=130)
        x += anchos[i]
    yy += 0.38
    _rect(s, M, yy - 0.06, W, 0.01, relleno=NEGRO)
    for fila in filas:
        # altura adaptativa: la fila crece con la celda que más se parta en líneas
        h = alto_fila
        for i, celda in enumerate(fila):
            txt = celda[0] if isinstance(celda, tuple) else celda
            t = celda[1].get("tam", tam) if isinstance(celda, tuple) else tam
            cpl = max(8, int((anchos[i] - 0.18) / (t * 0.0070)))
            n_lineas = max(1, -(-len(txt) // cpl))
            h = max(h, n_lineas * (t * 0.0182) + 0.19)
        x = M
        for i, celda in enumerate(fila):
            if isinstance(celda, tuple):
                txt, op = celda
            else:
                txt, op = celda, {}
            _tx(s, x, yy + 0.06, anchos[i] - 0.18, h - 0.06, txt,
                op.get("fuente", CUERPO), op.get("tam", tam),
                op.get("color", NEGRO), bold=op.get("bold", False), interlineado=1.2)
            x += anchos[i]
        yy += h
        _rect(s, M, yy - 0.04, W, 0.008, relleno=LINEA)
    _rect(s, M, yy + 0.02, W, 0.03, relleno=NEGRO)
    return yy + 0.05 + 0.30


def pasos(s, y, items, alto=0.78):
    """Lista numerada con número grande a la izquierda."""
    for i, (tit, desc) in enumerate(items):
        yy = y + i * alto
        _tx(s, M, yy, 0.62, 0.50, f"{i+1:02d}", TITULAR, 19, NARANJA)
        _tx(s, M + 0.78, yy + 0.02, W - 0.78, 0.26, tit, CUERPO, 12, NEGRO, bold=True)
        _tx(s, M + 0.78, yy + 0.28, W - 0.78, 0.44, desc, CUERPO, 11, TINTA, interlineado=1.25)
    return y + len(items) * alto + 0.20


def bloque_codigo(s, y, lineas, alto=None, titulo=None):
    h = alto or ((0.42 if titulo else 0.22) + 0.228 * len(lineas))
    _rect(s, M, y, W, h, relleno=RGBColor(0xF4, 0xF4, 0xF4), borde=None)
    if titulo:
        _tx(s, M + 0.22, y + 0.10, W - 0.44, 0.20, titulo, MONO, 7.3, GRIS_LBL, bold=True, spc=150)
        y0 = y + 0.36
    else:
        y0 = y + 0.13
    _tx(s, M + 0.22, y0, W - 0.44, h - (y0 - y) - 0.10, lineas, MONO, 9.5, NEGRO, interlineado=1.30)
    return y + h + 0.28


def nota(s, y, eyebrow, texto, alto=0.92):
    _rect(s, M, y, 0.05, alto, relleno=NARANJA)
    _tx(s, M + 0.26, y + 0.02, W - 0.30, 0.22, eyebrow, MONO, 8, NARANJA, bold=True, spc=190)
    _tx(s, M + 0.26, y + 0.28, W - 0.30, alto - 0.30, texto, CUERPO, 11.5, TINTA, interlineado=1.32)
    return y + alto + 0.28


def dos_columnas(s, y, izq, der, alto=2.3):
    """izq/der: (eyebrow, [parrafos]) sin marco, separadas por regla vertical."""
    ancho = (W - 0.60) / 2
    for i, (eb, parrafos) in enumerate([izq, der]):
        x = M + i * (ancho + 0.60)
        _tx(s, x, y, ancho, 0.22, eb, MONO, 8, NEGRO, bold=True, spc=190)
        _tx(s, x, y + 0.30, ancho, alto - 0.30, parrafos, CUERPO, 11.5, NEGRO,
            interlineado=1.28, espacio_post=7)
    _rect(s, M + ancho + 0.29, y, 0.01, alto, relleno=NEGRO)
    return y + alto + 0.28


# ---------- plantillas de cierre, replicadas de SIO0012-E1 (diapositivas 61, 62 y 64) ----------
GRIS_CAJA = RGBColor(0xF4, 0xF4, 0xF4)


def _cabecera(p, eyebrow, titulo, num, tam=26, alto=0.62):
    s = p.slides.add_slide(p.slide_layouts[6])
    _rect(s, M, 0.40, W, 0.05, relleno=NEGRO)
    _tx(s, M, 0.56, W, 0.22, eyebrow, MONO, 8.5, GRIS, bold=True, spc=240)
    _tx(s, M, 0.80, W, alto, titulo, TITULAR, tam, NEGRO, interlineado=1.08)
    _pie(p, s, num)
    return s


def glosario(p, eyebrow, parte, total, entradas, num,
             cierre="Toda sigla utilizada en la sesión aparece en el glosario con su forma completa en el "
                    "idioma original y su equivalencia en español."):
    """entradas: [(sigla, definición)] — hasta 10, cinco por columna, como en la referencia."""
    s = _cabecera(p, eyebrow, f"SIGLAS DE LA SESIÓN · {parte} DE {total}", num)
    for i, (sigla, defi) in enumerate(entradas[:10]):
        col, fila = i // 5, i % 5
        x = M + col * 6.24
        y = 1.62 + fila * 0.94
        _rect(s, x, y, 5.85, 0.84, relleno=GRIS_CAJA)
        _tx(s, x + 0.20, y + 0.10, 5.45, 0.26, sigla, TITULAR, 12.5, NEGRO)
        _tx(s, x + 0.20, y + 0.38, 5.45, 0.42, defi, CUERPO, 9.5, TINTA, interlineado=1.18)
    _tx(s, M, 6.36, W, 0.36, cierre, CUERPO, 11, TINTA)
    return s


def fuentes(p, eyebrow, titulo, refs, num, nota_eb=None, nota_txt=None):
    """refs: [(autor_año, título, apartado)] — hasta 7 por diapositiva."""
    s = _cabecera(p, eyebrow, titulo, num)
    paso = min(0.66, 3.90 / max(1, len(refs)))
    y = 1.72
    for autor, tit, apartado in refs:
        _tx(s, M, y, 2.55, paso - 0.10, autor, CUERPO, 10.5, NEGRO, bold=True, interlineado=1.15)
        _tx(s, 3.37, y, 5.55, paso - 0.10, tit, CUERPO, 10.5, NEGRO, interlineado=1.15)
        _tx(s, 9.17, y, 3.50, paso - 0.10, apartado, CUERPO, 10, TINTA, interlineado=1.15)
        _rect(s, M, y + paso - 0.10, W, 0.01, relleno=LINEA)
        y += paso
    if nota_eb:
        nota(s, 5.82, nota_eb, nota_txt, alto=0.86)
    return s


def trabajo_independiente(p, eyebrow, titulo_lineas, horas, lecturas, cond_eb, cond_txt, num,
                          titulo_lecturas=None):
    """horas: [(valor, l1, l2, destacado)] × 3 · lecturas: [(autor, qué leer y para qué)] hasta 3."""
    s = _cabecera(p, eyebrow, titulo_lineas, num, tam=23, alto=0.95)
    y = cifras(s, 1.96, horas)
    if titulo_lecturas:
        _tx(s, M, y - 0.12, W, 0.30, titulo_lecturas, CUERPO, 13.5, NEGRO, bold=True)
        y += 0.24
    for autor, que in lecturas[:3]:
        _tx(s, M, y, 2.90, 0.50, autor, CUERPO, 12, NEGRO, bold=True, interlineado=1.15)
        _tx(s, 3.72, y, 8.95, 0.50, que, CUERPO, 11.5, TINTA, interlineado=1.2)
        _rect(s, M, y + 0.54, W, 0.01, relleno=LINEA)
        y += 0.64
    nota(s, 5.82, cond_eb, cond_txt, alto=0.86)
    return s
