# -*- coding: utf-8 -*-
"""Detecta solapes entre cuadros de texto, desbordes y texto que invade la regla del pie."""
import sys
from pptx import Presentation
P = 914400

def rect(sh): return (sh.left/P, sh.top/P, (sh.left+sh.width)/P, (sh.top+sh.height)/P)

def area_comun(a, b):
    x = max(0, min(a[2], b[2]) - max(a[0], b[0]))
    y = max(0, min(a[3], b[3]) - max(a[1], b[1]))
    return x * y

def lineas(sh, par):
    txt = "".join(r.text for r in par.runs)
    tam = next((r.font.size.pt for r in par.runs if r.font.size), 11)
    cpl = max(6, int((sh.width / P) / (tam * 0.0070)))
    return txt, tam, max(1, -(-len(txt) // cpl))

def alto_estimado(sh):
    total = 0.0
    for par in sh.text_frame.paragraphs:
        txt, tam, n = lineas(sh, par)
        if txt.strip():
            total += n * tam * 0.0168 + (par.space_after.pt if par.space_after else 0) * 0.0139
    return total

def alto_real(sh):
    """Con el interlineado del párrafo y el espacio posterior, como lo dibuja el visor."""
    total = 0.0
    for par in sh.text_frame.paragraphs:
        txt, tam, n = lineas(sh, par)
        if txt.strip():
            ls = par.line_spacing if isinstance(par.line_spacing, float) else 1.0
            total += n * tam * 1.17 * ls / 72 + (par.space_after.pt if par.space_after else 0) / 72
    return total

pres = Presentation(sys.argv[1])
solapes = desbordes = pie = 0
for i, s in enumerate(pres.slides, 1):
    cajas = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    for j, a in enumerate(cajas):
        for b in cajas[j+1:]:
            ra, rb = rect(a), rect(b)
            com = area_comun(ra, rb)
            menor = min((ra[2]-ra[0])*(ra[3]-ra[1]), (rb[2]-rb[0])*(rb[3]-rb[1]))
            if menor > 0 and com / menor > 0.28:
                print(f"  SOLAPE   diapositiva {i:2}: '{a.text_frame.text[:34]}' × '{b.text_frame.text[:34]}'")
                solapes += 1
        h = alto_estimado(a)
        if h > (a.height/P) * 1.35 and h > 0.30:
            print(f"  DESBORDE diapositiva {i:2}: '{a.text_frame.text[:46]}' necesita {h:.2f}\" en {a.height/P:.2f}\"")
            desbordes += 1
        top = a.top / P
        if top < 6.78 and "SIO0010 · " not in a.text_frame.text:
            fondo = top + alto_real(a)
            if fondo > 6.78:
                print(f"  PIE      diapositiva {i:2}: '{a.text_frame.text[:46]}' llega a {fondo:.2f}\"")
                pie += 1
print(f"\nSolapes: {solapes} · Desbordes: {desbordes} · Invaden el pie: {pie}")
